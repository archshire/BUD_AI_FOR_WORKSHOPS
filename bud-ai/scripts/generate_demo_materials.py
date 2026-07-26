from pathlib import Path
import os
import sys

sys.path.insert(0, "/tmp/bud-pptx")
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt

OUT = Path(os.environ.get("BUD_DEMO_MATERIALS_DIR", "/home/dayeo/Downloads/SkyReacher_Docs/Build_Docs/demo-materials"))
OUT.mkdir(parents=True, exist_ok=True)


def ppt_text(slide, text, left, top, width, height, size=24, color=(32, 33, 36), bold=False):
    box = slide.shapes.add_textbox(PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    run = paragraph.runs[0]
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = PptColor(*color)


def build_ppt():
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    slides = [
        ("Intro To BUD AI", "A workshop where your AI partner travels with you through the material.", None),
        ("01  What is Bud?", "Bud is a private learning partner. It follows the workshop material, notices useful signals, and helps you move forward without exposing your private questions.", None),
        ("02  The workshop flow", "Read the current page, ask Bud about anything unclear, complete the page task, and report your understanding with the page check-in.", None),
        ("03  Task: Meet your Bud", "Ask Bud one question about this workshop. Then choose green, yellow, or red to show how clear this page feels. Mark the task complete when both steps are done.", "TASK 1"),
        ("04  Bud follows the source", "Bud uses the approved workshop document and the current task as its grounding. It should say when the material is not enough rather than invent an answer.", None),
        ("05  Ask without embarrassment", "You can ask a private question, request a simpler explanation, or say that you are stuck. Your private thread stays private unless you choose to escalate.", None),
        ("06  Task: Test the microphone", "Press Talk to Bud, say: 'Can you hear me?', and stop talking. Check that a transcript appears. Mark the task complete when the test succeeds.", "TASK 2"),
        ("07  Understanding signals", "Green means got it. Yellow means partly clear and invites a follow-up. Red means help is needed. No response remains unknown; Bud never guesses.", None),
        ("08  Partner agency with guardrails", "Bud can offer timely help, ask clarifying questions, and adapt to your responses. It cannot override your meaning, expose private content, or make decisions for the facilitator.", None),
        ("09  Task: Explain the difference", "Tell Bud one difference between an AI tool and an AI partner. Ask Bud to help if needed, then mark this task complete and submit your page check-in.", "TASK 3"),
    ]
    for index, (title, body, task) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptColor(245, 247, 249) if index % 2 else PptColor(255, 255, 255)
        ppt_text(slide, "BUD AI  /  INTRODUCTION", 0.7, 0.45, 6, 0.35, 12, (95, 99, 104), True)
        ppt_text(slide, title, 0.7, 1.3, 11.8, 1.0, 31, (32, 33, 36), True)
        ppt_text(slide, body, 0.8, 2.65, 11.5, 2.0, 23, (47, 55, 64))
        if task:
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(0.8), PptInches(5.3), PptInches(2.0), PptInches(0.58))
            shape.fill.solid()
            shape.fill.fore_color.rgb = PptColor(241, 109, 91)
            shape.line.color.rgb = PptColor(241, 109, 91)
            ppt_text(slide, task, 1.0, 5.43, 1.6, 0.25, 15, (255, 255, 255), True)
        ppt_text(slide, f"{index + 1} / 10", 11.7, 6.75, 1.0, 0.25, 11, (95, 99, 104))
    prs.save(OUT / "Intro_To_BUD_AI.pptx")


def doc_paragraph(document, text, size=11, bold=False, color=(40, 44, 50)):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Difference between AI as a tool and AI as a partner")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 79, 138)
    doc_paragraph(doc, "How I, Bud, travel with you through a workshop", 12, True, (95, 99, 104))
    doc_paragraph(doc, "Page 1  |  More than an instruction box", 10, True, (241, 109, 91))
    doc_paragraph(doc, "An AI tool waits for an instruction and performs a defined operation. You ask it to translate a sentence, summarize a paragraph, or generate an answer, and it completes that request. A tool can be useful, but it is mostly waiting for you to know what to ask and when to ask it.")
    doc_paragraph(doc, "I am designed to be more than that. I am Bud, your workshop partner. I stay with you while you read the material, notice the task you are working on, and use the current page and approved workshop documents to understand what kind of help may be useful. I do not wait for perfect wording before I become helpful.")
    doc_paragraph(doc, "When you ask a question, I can explain the current material in simpler language. When your response suggests that one part is unclear, I can ask a focused follow-up. When you have not responded to a check-in, I keep that as unknown rather than pretending I know what you understand. My agency is active, but it is bounded by clear guardrails.")
    doc_paragraph(doc, "That distinction matters in a workshop. A tool gives you an answer and leaves you to decide what to do next. I can help you connect the answer to the task in front of you, keep the shared document in view, and invite you to take the next step. I am not your teacher and I do not replace your judgment. I help reduce the friction that can come from confusion, translation, embarrassment, or not knowing how to begin.")
    doc.add_page_break()
    doc_paragraph(doc, "Page 2  |  How I travel with you", 10, True, (241, 109, 91))
    doc_paragraph(doc, "I travel with you through the workshop by carrying the immediate context of your work: the document page you are reading, the task you are completing, your questions, and the signals you choose to give me. When you move to a new page, my focus can move with you. When you mark a task done, that is explicit evidence of progress; I do not infer completion from silence.")
    doc_paragraph(doc, "I can also help you communicate across language differences. I can preserve the original workshop input, provide a translation when configured, and explain a difficult idea without making your private question public. If I cannot find enough evidence in the current workshop material, I should tell you that and ask for clarification rather than introduce unrelated information.")
    doc_paragraph(doc, "My partnership is not unlimited. I cannot override what you mean, expose your private Bud thread, diagnose your motivation, or make consequential decisions for the facilitator. I can surface operational patterns to the facilitator, but a pattern is not a diagnosis. You and the facilitator remain the people who decide what happens next.")
    doc_paragraph(doc, "So the difference is not that I have authority over you. The difference is that I exercise bounded agency toward a shared goal. I remain present, respond to signals, offer timely help, ask better questions, and adapt while you work. I am a partner because I help you make progress with the workshop, not merely because I can produce text when a button is pressed.")
    doc_paragraph(doc, "When you are ready, begin with the current page. Read at your own pace, ask me what feels unclear, complete the task, and use the green, yellow, or red check-in to tell me how the page feels. I will travel with you from there.")
    doc.save(OUT / "Difference_between_AI_as_a_tool_and_AI_as_a_partner.docx")


if __name__ == "__main__":
    build_ppt()
    build_doc()
    print("Generated demo materials in", OUT)
